#!/usr/bin/env python3
"""Build restrained DOCX instruments from reviewed JSON. Render separately."""
import argparse,json,re,hashlib
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def validate(data):
    for k in ('id','title','status','filing_gate','sections','sources'):
        if k not in data: raise ValueError('Missing '+k)
    if data['status'] not in ('prepared','conditional','internal'): raise ValueError('Invalid status')
    if not re.fullmatch(r'[A-Za-z0-9_]+',data['id']): raise ValueError('Unsafe document id')
    if not data['sections'] or not data['sources']: raise ValueError('Content and sources required')
    for s in data['sections']:
        if not isinstance(s.get('paragraphs'),list) or not all(isinstance(p,str) for p in s['paragraphs']): raise ValueError('Invalid paragraphs')
    for s in data['sources']:
        if not s.get('url','').startswith('https://'): raise ValueError('Source URL missing')

def make(data,out):
    validate(data)
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.82)
    sec.left_margin=sec.right_margin=Inches(.94)
    sec.header_distance=Inches(.3); sec.footer_distance=Inches(.4)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(11)
        rf=s.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme'):
            rf.attrib.pop(qn('w:'+attr),None)
        for attr in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+attr),'Times New Roman')
        s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(6)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Normal'].paragraph_format.widow_control=True
    d.styles['Title'].font.size=Pt(14);d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(13)
    d.styles['Heading 1'].font.size=Pt(11);d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    d.styles['Heading 1'].paragraph_format.keep_with_next=True
    for node in list(d.styles.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
    d.add_paragraph(data['title'].replace(',',''),style='Title')
    for s in data['sections']:
        if s.get('heading'):d.add_paragraph(re.sub(r'^(\d+)\.\s*',r'\1 ',s['heading']).replace(',',''),style='Heading 1')
        for p in s['paragraphs']:
            x=d.add_paragraph(p);x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=f.add_run();r.font.name='Times New Roman';r.font.size=Pt(10)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    d.core_properties.title=data['title'];d.core_properties.subject=data['id'];d.core_properties.author=''
    suffix='_CONDITIONAL' if data['status']=='conditional' else ''
    path=out/(data['id']+suffix+'.docx');d.save(path)
    return path

def main():
    ap=argparse.ArgumentParser();ap.add_argument('input');ap.add_argument('output');a=ap.parse_args()
    source=Path(a.input);out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    inputs=sorted(source.glob('*.json')) if source.is_dir() else [source]
    ids=set()
    for p in inputs:
        data=json.loads(p.read_text());validate(data)
        if data['id'] in ids:raise ValueError('Duplicate id')
        ids.add(data['id']);dest=make(data,out);print(dest)
if __name__=='__main__':main()
